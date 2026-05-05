import logging
from pathlib import Path
from typing import Any, Dict, List, Optional, Union

import pandas as pd
from openpyxl import load_workbook

logger = logging.getLogger(__name__)


class ExcelUtils:
    """
    Excel utility class.
    Supports reading, writing, DB-style querying, and data validation of Excel files.
    """

    def __init__(self, file_path: Union[str, Path]):
        self.file_path = Path(file_path)
        if not self.file_path.exists():
            raise FileNotFoundError(f"Excel file not found: {self.file_path}")
        logger.info("ExcelUtils ready: %s", self.file_path)

    # ------------------------------------------------------------------
    # Readers
    # ------------------------------------------------------------------

    def read_as_dataframe(self, sheet_name: Union[str, int] = 0) -> pd.DataFrame:
        try:
            df = pd.read_excel(self.file_path, sheet_name=sheet_name)
            logger.info("Read %d rows from sheet '%s'", len(df), sheet_name)
            return df
        except Exception as exc:
            logger.error("Failed to read sheet '%s': %s", sheet_name, exc)
            raise

    def read_as_dict_list(self, sheet_name: Union[str, int] = 0) -> List[Dict]:
        return self.read_as_dataframe(sheet_name).to_dict(orient="records")

    def get_cell_value(self, row: int, col: int, sheet_name: str = None) -> Any:
        try:
            wb = load_workbook(self.file_path, data_only=True)
            ws = wb[sheet_name] if sheet_name else wb.active
            value = ws.cell(row=row, column=col).value
            wb.close()
            logger.info("Cell (%d, %d) = %s", row, col, value)
            return value
        except Exception as exc:
            logger.error("get_cell_value failed: %s", exc)
            raise

    def get_column_values(
        self, column: Union[str, int], sheet_name: Union[str, int] = 0
    ) -> List[Any]:
        df = self.read_as_dataframe(sheet_name)
        if isinstance(column, int):
            return df.iloc[:, column].tolist()
        return df[column].tolist()

    def get_sheet_names(self) -> List[str]:
        wb = load_workbook(self.file_path, read_only=True)
        names = wb.sheetnames
        wb.close()
        return names

    def get_column_names(self, sheet_name: Union[str, int] = 0) -> List[str]:
        return list(self.read_as_dataframe(sheet_name).columns)

    def get_row_count(self, sheet_name: Union[str, int] = 0) -> int:
        return len(self.read_as_dataframe(sheet_name))

    # ------------------------------------------------------------------
    # DB-style querying
    # ------------------------------------------------------------------

    def find_row(
        self, column: str, value: Any, sheet_name: Union[str, int] = 0
    ) -> Optional[Dict]:
        df = self.read_as_dataframe(sheet_name)
        result = df[df[column] == value]
        if result.empty:
            logger.warning("No row found where %s == %s", column, value)
            return None
        return result.iloc[0].to_dict()

    def find_all_rows(
        self, column: str, value: Any, sheet_name: Union[str, int] = 0
    ) -> List[Dict]:
        df = self.read_as_dataframe(sheet_name)
        return df[df[column] == value].to_dict(orient="records")

    def query(self, condition: str, sheet_name: Union[str, int] = 0) -> List[Dict]:
        """
        Query rows using pandas query syntax.
        Example: query("age > 30 and city == 'New York'")
        """
        try:
            df = self.read_as_dataframe(sheet_name)
            result = df.query(condition)
            logger.info("Query '%s' → %d row(s)", condition, len(result))
            return result.to_dict(orient="records")
        except Exception as exc:
            logger.error("Query failed: %s", exc)
            raise

    def group_by(
        self, column: str, sheet_name: Union[str, int] = 0
    ) -> Dict[Any, List[Dict]]:
        df = self.read_as_dataframe(sheet_name)
        groups: Dict[Any, List[Dict]] = {}
        for key, grp in df.groupby(column):
            groups[key] = grp.to_dict(orient="records")
        return groups

    def get_unique_values(
        self, column: str, sheet_name: Union[str, int] = 0
    ) -> List[Any]:
        df = self.read_as_dataframe(sheet_name)
        return sorted(df[column].dropna().unique().tolist())

    # ------------------------------------------------------------------
    # Writers
    # ------------------------------------------------------------------

    def write_data(
        self,
        data: List[Dict],
        sheet_name: str = "Sheet1",
        output_path: Union[str, Path] = None,
    ) -> None:
        output = Path(output_path) if output_path else self.file_path
        df = pd.DataFrame(data)
        mode = "a" if output.exists() else "w"
        with pd.ExcelWriter(output, engine="openpyxl", mode=mode) as writer:
            df.to_excel(writer, sheet_name=sheet_name, index=False)
        logger.info("Wrote %d rows to '%s' in %s", len(data), sheet_name, output)

    @staticmethod
    def create_excel(
        file_path: Union[str, Path], sheets_data: Dict[str, List[Dict]]
    ) -> None:
        path = Path(file_path)
        path.parent.mkdir(parents=True, exist_ok=True)
        with pd.ExcelWriter(path, engine="openpyxl") as writer:
            for sheet_name, rows in sheets_data.items():
                pd.DataFrame(rows).to_excel(writer, sheet_name=sheet_name, index=False)
        logger.info("Created Excel: %s", path)

    # ------------------------------------------------------------------
    # Validators
    # ------------------------------------------------------------------

    def validate_no_nulls(
        self, column: str, sheet_name: Union[str, int] = 0
    ) -> bool:
        df = self.read_as_dataframe(sheet_name)
        nulls = df[column].isnull().sum()
        assert nulls == 0, f"Column '{column}' has {nulls} null value(s)"
        logger.info("[PASS] Column '%s' has no nulls", column)
        return True

    def validate_unique(
        self, column: str, sheet_name: Union[str, int] = 0
    ) -> bool:
        df = self.read_as_dataframe(sheet_name)
        duplicates = df[column][df[column].duplicated()].tolist()
        assert not duplicates, f"Column '{column}' has duplicates: {duplicates}"
        logger.info("[PASS] Column '%s' values are unique", column)
        return True

    def validate_column_dtype(
        self, column: str, expected_dtype: str, sheet_name: Union[str, int] = 0
    ) -> bool:
        df = self.read_as_dataframe(sheet_name)
        actual = str(df[column].dtype)
        assert expected_dtype in actual, (
            f"Column '{column}': expected dtype '{expected_dtype}', got '{actual}'"
        )
        logger.info("[PASS] Column '%s' dtype '%s' contains '%s'", column, actual, expected_dtype)
        return True

    def validate_row_count(
        self, expected_count: int, sheet_name: Union[str, int] = 0
    ) -> bool:
        actual = self.get_row_count(sheet_name)
        assert actual == expected_count, f"Expected {expected_count} rows, got {actual}"
        logger.info("[PASS] Row count is %d", actual)
        return True

    def validate_allowed_values(
        self,
        column: str,
        allowed: List[Any],
        sheet_name: Union[str, int] = 0,
    ) -> bool:
        df = self.read_as_dataframe(sheet_name)
        invalid = df[~df[column].isin(allowed)][column].tolist()
        assert not invalid, (
            f"Column '{column}' has invalid values: {invalid}. Allowed: {allowed}"
        )
        logger.info("[PASS] All values in '%s' are allowed", column)
        return True

    def validate_numeric_range(
        self,
        column: str,
        min_val: Union[int, float],
        max_val: Union[int, float],
        sheet_name: Union[str, int] = 0,
    ) -> bool:
        df = self.read_as_dataframe(sheet_name)
        out_of_range = df[(df[column] < min_val) | (df[column] > max_val)][column].tolist()
        assert not out_of_range, (
            f"Column '{column}' has values outside [{min_val}, {max_val}]: {out_of_range}"
        )
        logger.info("[PASS] All values in '%s' are within [%s, %s]", column, min_val, max_val)
        return True

    def validate_columns_exist(
        self, required_columns: List[str], sheet_name: Union[str, int] = 0
    ) -> bool:
        actual_cols = self.get_column_names(sheet_name)
        missing = [c for c in required_columns if c not in actual_cols]
        assert not missing, f"Missing columns: {missing}. Available: {actual_cols}"
        logger.info("[PASS] All required columns present")
        return True
